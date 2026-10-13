from io import BytesIO
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from .skill_extractor import extract_skills


JOBS = [
    {
        "title": "Python Developer",
        "company": "Tech Solutions",
        "skills": ["python", "django", "sql", "git"],
        "description": (
            "Build and maintain Python backend applications and APIs. Use Python, "
            "Django, SQL databases, and Git to deliver reliable software. Work with "
            "product and engineering teams, write tests, debug issues, and document "
            "solutions. Relevant experience may include backend development, API "
            "design, database integration, and software projects."
        ),
    },
    {
        "title": "AI/ML Engineer",
        "company": "AI Technologies",
        "skills": ["python", "machine learning", "tensorflow", "pandas", "numpy"],
        "description": (
            "Develop and evaluate machine-learning models with Python. Prepare data "
            "using pandas and NumPy, apply machine-learning methods, and work with "
            "frameworks such as TensorFlow. Communicate results and document model "
            "experiments. Relevant experience may include data preparation, model "
            "training, evaluation, and AI projects."
        ),
    },
    {
        "title": "Frontend Developer",
        "company": "Web Solutions",
        "skills": ["html", "css", "javascript", "react"],
        "description": (
            "Build responsive web interfaces using HTML, CSS, JavaScript, and React. "
            "Create accessible components, connect user interfaces to APIs, and test "
            "across browsers. Collaborate with design and backend teams. Relevant "
            "experience may include frontend applications, responsive design, and "
            "web development projects."
        ),
    },
    {
        "title": "Data Analyst",
        "company": "Data Corp",
        "skills": ["python", "sql", "excel", "power bi", "pandas"],
        "description": (
            "Analyze and report on data using SQL, Python, Excel, pandas, and Power BI. "
            "Prepare datasets, identify trends, build dashboards, and explain findings "
            "to stakeholders. Relevant experience may include reporting, data "
            "visualization, and analytical projects."
        ),
    },
    {
        "title": "Software Developer",
        "company": "Sample role catalog",
        "skills": ["python", "java", "git", "sql", "software testing"],
        "description": (
            "Design, develop, test, and maintain software applications. Use one or "
            "more programming languages such as Python or Java, collaborate with "
            "engineering teams, use Git for version control, work with SQL databases, "
            "and write or run software tests. Relevant experience may include "
            "application development, debugging, code review, and software projects."
        ),
    },
    {
        "title": "Backend Developer",
        "company": "Sample role catalog",
        "skills": ["python", "fastapi", "django", "sql", "rest api", "git"],
        "description": (
            "Build and maintain server-side applications, REST APIs, and database "
            "integrations. Use Python frameworks such as FastAPI or Django, SQL, and "
            "Git. Write tests, troubleshoot production issues, and consider security, "
            "performance, and reliability. Relevant experience may include API "
            "development and backend services."
        ),
    },
    {
        "title": "Full Stack Developer",
        "company": "Sample role catalog",
        "skills": ["html", "css", "javascript", "react", "node.js", "sql"],
        "description": (
            "Develop end-to-end web applications across frontend and backend layers. "
            "Build accessible interfaces with HTML, CSS, JavaScript, and React; "
            "develop server-side features with Node.js; and connect applications "
            "to SQL databases and APIs. Collaborate with design and engineering "
            "teams, test features, and maintain web applications."
        ),
    },
    {
        "title": "Financial Analyst",
        "company": "Sample role catalog",
        "skills": ["excel", "financial analysis", "financial modeling", "forecasting", "budgeting"],
        "description": (
            "Support business decisions through financial analysis, budgeting, and "
            "forecasting. Build and maintain spreadsheet-based financial models, "
            "review financial statements and performance data, investigate variances, "
            "and present clear findings to stakeholders. Relevant experience may "
            "include Excel analysis, planning, reporting, and finance projects."
        ),
    },
]

MATCH_THRESHOLD = 50
RESUME_SECTION_HEADINGS = {
    "summary",
    "professional summary",
    "profile",
    "objective",
    "career objective",
    "skills",
    "technical skills",
    "core competencies",
    "work experience",
    "professional experience",
    "employment history",
    "experience",
    "education",
    "academic background",
    "projects",
    "certifications",
    "certificates",
}
RESUME_SECTION_ALIASES = {
    "Summary": {"summary", "professional summary", "profile", "objective", "career objective"},
    "Skills": {"skills", "technical skills", "core competencies"},
    "Experience": {"experience", "work experience", "professional experience", "employment history"},
    "Education": {"education", "academic background"},
    "Projects": {"projects", "project experience"},
    "Certifications": {"certifications", "certificates", "licenses"},
}


def match_jobs(user_skills):
    if isinstance(user_skills, str):
        user_skill_set = set(extract_skills(user_skills))
    else:
        user_skill_set = set(user_skills)

    results = []

    for job in JOBS:
        required = set(job["skills"])
        matched = required & user_skill_set
        score = (len(matched) / len(required)) * 100 if required else 0
        result = job.copy()
        result["matched_skills"] = sorted(matched)
        result["match_score"] = round(score, 2)
        results.append(result)

    results.sort(key=lambda x: x["match_score"], reverse=True)
    return results


def get_job_description(title):
    return next((job for job in JOBS if job["title"] == title), None)


def review_resume_format(resume_text):
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    normalized_headings = {
        re.sub(r"[:\s]+$", "", line.lower()).strip()
        for line in lines
    }
    sections_found = [
        section
        for section, aliases in RESUME_SECTION_ALIASES.items()
        if normalized_headings.intersection(aliases)
    ]
    issues = []

    if "Experience" not in sections_found:
        issues.append("Add a clearly labeled Experience section, or state that you are an entry-level candidate.")
    if "Education" not in sections_found:
        issues.append("Add a clearly labeled Education section.")
    if "Skills" not in sections_found:
        issues.append("Add a clearly labeled Skills section to make relevant abilities easy to scan.")
    if "Projects" not in sections_found:
        issues.append("Add a Projects section if you have relevant coursework, personal, or professional projects.")
    has_email = re.search(r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b", resume_text)
    has_phone = re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", resume_text)
    if not has_email or not has_phone:
        issues.append("Check that your contact details include both an email address and a phone number.")
    if not any(line.startswith(("-", "•", "*")) for line in lines):
        issues.append("Use concise bullet points for experience and project achievements.")
    if len(lines) < 8:
        issues.append("The extracted resume is very short; check that all pages and sections were included.")

    return {
        "issues": issues,
        "sections_found": sections_found,
        "note": (
            "This automated review checks extracted text structure only. "
            "It cannot judge visual layout, spelling, or whether your claims are accurate."
        ),
    }


def build_tailored_resume_docx(job, resume_text, resume_skills):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(40, 48, 64)

    title = document.add_heading("RESUME", level=0)
    title.alignment = 1
    title.runs[0].font.color.rgb = RGBColor(24, 40, 72)
    title.runs[0].font.size = Pt(23)

    role_paragraph = document.add_paragraph()
    role_paragraph.alignment = 1
    role_run = role_paragraph.add_run(f"TARGET ROLE: {job['title']}")
    role_run.bold = True
    role_run.font.color.rgb = RGBColor(49, 94, 251)

    detected_skills = ", ".join(resume_skills) or "No known skills were detected"
    skills_heading = document.add_heading("SKILLS FOUND IN YOUR RESUME", level=1)
    skills_heading.runs[0].font.color.rgb = RGBColor(24, 40, 72)
    document.add_paragraph(detected_skills)

    missing_role_skills = [
        skill for skill in job["skills"] if skill not in set(resume_skills)
    ]
    gap_heading = document.add_heading("ROLE SKILLS TO CHECK", level=1)
    gap_heading.runs[0].font.color.rgb = RGBColor(24, 40, 72)
    if missing_role_skills:
        document.add_paragraph(
            "These skills were not detected in the uploaded resume. Add them only if you have the relevant experience:"
        )
        for skill in missing_role_skills:
            document.add_paragraph(skill, style="List Bullet")
    else:
        document.add_paragraph("The listed role skills were detected in your resume.")

    content_heading = document.add_heading("YOUR RESUME CONTENT", level=1)
    content_heading.runs[0].font.color.rgb = RGBColor(24, 40, 72)

    for raw_line in resume_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        heading_key = line.strip(":").lower()
        if heading_key in RESUME_SECTION_HEADINGS and len(line) <= 50:
            heading = document.add_heading(line.strip(":"), level=2)
            heading.runs[0].font.color.rgb = RGBColor(24, 40, 72)
        elif line.startswith(("-", "•", "*")):
            document.add_paragraph(line.lstrip("-•* ").strip(), style="List Bullet")
        else:
            document.add_paragraph(line)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note_run = note.add_run(
        "Review before applying: this document reformats the extracted text from your upload. "
        "Verify every detail, correct extraction errors, and add role-relevant skills only if you genuinely have them."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(91, 100, 117)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def build_sample_resume(job, resume_skills):
    role_skills = ", ".join(job["skills"])
    candidate_skills = ", ".join(resume_skills) or "[Add your relevant skills]"

    return f"""[YOUR NAME]
[City, Country] | [Phone number] | [Email address] | [LinkedIn profile]

TARGET ROLE
{job["title"]}

PROFESSIONAL SUMMARY
Professional seeking a {job["title"]} role, with experience in [your relevant experience or area of study]. Skilled in {candidate_skills}. Add a short, truthful statement about the work you have done and the value you can bring.

TECHNICAL SKILLS
Skills identified in your resume: {candidate_skills}
Other skills relevant to this role (include only those you can demonstrate): {role_skills}

PROJECTS
[Relevant project name] | [Dates]
- Describe what you personally built, analyzed, or improved using [relevant tools or skills].
- Add a measurable result if available, such as [time saved, accuracy, users, or data volume].

EXPERIENCE
[Job title, internship, or volunteer role] | [Organization] | [Dates]
- Describe a responsibility relevant to {job["title"]} and the outcome you achieved.
- Include specific tools, scope, and measurable results only when accurate.

EDUCATION
[Degree or qualification] | [Institution] | [Graduation year]

NOTE
This is an illustrative template, not a record of your experience. Replace every bracketed placeholder and include only skills, projects, and achievements that are true for you.
"""


def build_sample_resume_docx(job, resume_skills):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(40, 48, 64)

    title = document.add_heading("[YOUR NAME]", level=0)
    title.alignment = 1
    title.runs[0].font.color.rgb = RGBColor(24, 40, 72)
    title.runs[0].font.size = Pt(24)

    contact = document.add_paragraph(
        "[City, Country] | [Phone number] | [Email address] | [LinkedIn profile]"
    )
    contact.alignment = 1

    target = document.add_paragraph()
    target.alignment = 1
    target_run = target.add_run(f"TARGET ROLE: {job['title']}")
    target_run.bold = True
    target_run.font.color.rgb = RGBColor(49, 94, 251)

    candidate_skills = ", ".join(resume_skills) or "[Add your relevant skills]"
    suggested_skills = ", ".join(job["skills"])

    sections = [
        (
            "PROFESSIONAL SUMMARY",
            [
                f"Professional seeking a {job['title']} role, with experience in [your relevant experience or area of study]. "
                f"Skilled in {candidate_skills}. Add a short, truthful statement about the work you have done and the value you can bring."
            ],
        ),
        (
            "TECHNICAL SKILLS",
            [
                f"Skills identified in your resume: {candidate_skills}",
                f"Other skills relevant to this role (include only those you can demonstrate): {suggested_skills}",
            ],
        ),
        (
            "PROJECTS",
            [
                "[Relevant project name] | [Dates]",
                "Describe what you personally built, analyzed, or improved using relevant tools or skills.",
                "Add a measurable result if available, such as time saved, accuracy, users, or data volume.",
            ],
        ),
        (
            "EXPERIENCE",
            [
                "[Job title, internship, or volunteer role] | [Organization] | [Dates]",
                f"Describe a responsibility relevant to {job['title']} and the outcome you achieved.",
                "Include specific tools, scope, and measurable results only when accurate.",
            ],
        ),
        (
            "EDUCATION",
            ["[Degree or qualification] | [Institution] | [Graduation year]"],
        ),
        (
            "CERTIFICATIONS",
            ["[Certification name] | [Issuing organization] | [Year, if applicable]"],
        ),
    ]

    for heading, content in sections:
        section_heading = document.add_heading(heading, level=1)
        section_heading.runs[0].font.color.rgb = RGBColor(24, 40, 72)
        for index, text in enumerate(content):
            if heading in {"PROJECTS", "EXPERIENCE"} and index > 0:
                document.add_paragraph(text, style="List Bullet")
            else:
                paragraph = document.add_paragraph(text)
                if heading in {"PROJECTS", "EXPERIENCE"} and index == 0:
                    paragraph.runs[0].bold = True

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note_run = note.add_run(
        "TEMPLATE NOTE: This is an illustrative template, not a record of your experience. "
        "Replace placeholders and include only details that are true for you."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(91, 100, 117)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output